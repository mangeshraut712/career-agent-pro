"""Resume parser for extracting text from various file formats."""
import io
import logging

import PyPDF2
import docx

from security_utils import MAX_EXTRACTED_CHARS, MAX_PDF_PAGES, ResumeUploadError, validate_resume_filename

logger = logging.getLogger(__name__)


class ResumeParser:
    """Parser for extracting text from resume files."""
    
    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> str:
        """Extract text from PDF file."""
        text = ""
        try:
            reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
            page_count = len(reader.pages)
            if page_count == 0:
                logger.warning("PDF file has no pages")
                return text
            if page_count > MAX_PDF_PAGES:
                raise ResumeUploadError(f"Resume PDF exceeds the {MAX_PDF_PAGES} page limit.")
            
            for page_num, page in enumerate(reader.pages):
                try:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                except Exception as e:
                    logger.warning(f"Error extracting text from page {page_num + 1}: {str(e)}")
                    continue
            
            if not text.strip():
                logger.warning("No text extracted from PDF")
        except ResumeUploadError:
            raise
        except PyPDF2.errors.PdfReadError:
            logger.exception("PDF read error")
            raise ResumeUploadError("Could not read file. Supported formats: PDF, DOCX, TXT.")
        except Exception:
            logger.exception("PDF extraction error")
            raise ResumeUploadError("Could not read file. Supported formats: PDF, DOCX, TXT.")
        
        return text.strip()

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        """Extract text from DOCX file."""
        text = ""
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            
            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    text += para.text + "\n"
            
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        text += row_text + "\n"
            
            if not text.strip():
                logger.warning("No text extracted from DOCX")
        except ResumeUploadError:
            raise
        except Exception:
            logger.exception("DOCX extraction error")
            raise ResumeUploadError("Could not read file. Supported formats: PDF, DOCX, TXT.")
        
        return text.strip()

    @classmethod
    def extract_text(cls, file_bytes: bytes, filename: str) -> str:
        """Extract text from file based on extension."""
        if not filename:
            raise ValueError("Filename is required")
        
        try:
            extension = validate_resume_filename(filename)
            if extension == ".pdf":
                text = cls.extract_text_from_pdf(file_bytes)
            elif extension == ".docx":
                text = cls.extract_text_from_docx(file_bytes)
            else:
                try:
                    text = file_bytes.decode("utf-8")
                except UnicodeDecodeError:
                    text = file_bytes.decode("latin-1", errors="ignore")
        except ResumeUploadError:
            raise
        except ValueError:
            raise ResumeUploadError("Could not read file. Supported formats: PDF, DOCX, TXT.")
        except Exception:
            logger.exception("Unexpected error extracting text")
            raise ResumeUploadError("Could not read file. Supported formats: PDF, DOCX, TXT.")

        if len(text) > MAX_EXTRACTED_CHARS:
            return text[:MAX_EXTRACTED_CHARS]
        return text

