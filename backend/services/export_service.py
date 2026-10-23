"""Export service for generating resumes in various formats."""
import logging
import os
from datetime import datetime
from typing import Optional
from docx import Document
from fpdf import FPDF
from schemas import ResumeData
from slugify import slugify

logger = logging.getLogger(__name__)

_LATEX_ESCAPES = {
    "\\": r"\textbackslash{}",
    "&": r"\&",
    "%": r"\%",
    "$": r"\$",
    "#": r"\#",
    "_": r"\_",
    "{": r"\{",
    "}": r"\}",
    "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}",
}

_DEJAVU_REGULAR = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
_DEJAVU_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
_DEJAVU_OBLIQUE = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"


def escape_latex(text: object) -> str:
    """Escape LaTeX specials in one pass so replacements are not reprocessed."""
    if text is None:
        return ""
    return "".join(_LATEX_ESCAPES.get(char, char) for char in str(text))


def _register_pdf_font(pdf: FPDF) -> str:
    """Use a Unicode font when it is installed. Otherwise fall back to Helvetica."""
    if not os.path.exists(_DEJAVU_REGULAR):
        return "Helvetica"
    try:
        pdf.add_font("DejaVu", "", _DEJAVU_REGULAR)
        if os.path.exists(_DEJAVU_BOLD):
            pdf.add_font("DejaVu", "B", _DEJAVU_BOLD)
        if os.path.exists(_DEJAVU_OBLIQUE):
            pdf.add_font("DejaVu", "I", _DEJAVU_OBLIQUE)
        return "DejaVu"
    except Exception:
        logger.exception("Failed to register Unicode PDF font")
        return "Helvetica"


def pdf_text(text: object, unicode_font: bool) -> str:
    value = "" if text is None else str(text)
    if unicode_font:
        return value
    return value.encode("latin-1", errors="replace").decode("latin-1")


class ExportService:
    """Service for exporting resumes to various formats."""
    
    @staticmethod
    def get_organized_path(job_title: Optional[str] = None, base_dir: str = "storage/exports") -> str:
        """Creates an organized path: storage/exports/YYYY-MM-DD/job-title/"""
        try:
            date_str = datetime.now().strftime("%Y-%m-%d")
            slug_title = slugify(job_title or "general")
            path = os.path.join(base_dir, date_str, slug_title)
            os.makedirs(path, exist_ok=True)
            return path
        except Exception as e:
            logger.error(f"Failed to create organized path: {str(e)}")
            raise
    
    @staticmethod
    def to_docx(resume: ResumeData, output_path: str) -> str:
        """Export resume to DOCX format."""
        try:
            if not resume.name:
                raise ValueError("Resume name is required")
            
            doc = Document()
            doc.add_heading(resume.name, 0)
            
            # Contact information
            contact_parts = [resume.email]
            if resume.phone:
                contact_parts.append(resume.phone)
            if resume.linkedin:
                contact_parts.append(resume.linkedin)
            
            if contact_parts:
                p = doc.add_paragraph()
                p.add_run(" | ".join(contact_parts)).italic = True
            
            # Summary
            if resume.summary:
                doc.add_heading('Summary', level=1)
                doc.add_paragraph(resume.summary)
            
            # Experience
            if resume.experience:
                doc.add_heading('Experience', level=1)
                for exp in resume.experience:
                    if not isinstance(exp, dict):
                        continue
                    role = exp.get('role', 'N/A')
                    company = exp.get('company', 'N/A')
                    doc.add_heading(f"{role} at {company}", level=2)
                    
                    if exp.get('duration'):
                        doc.add_paragraph(exp['duration']).italic = True
                    
                    if exp.get('description'):
                        doc.add_paragraph(exp['description'])
            
            # Education
            if resume.education:
                doc.add_heading('Education', level=1)
                for edu in resume.education:
                    if not isinstance(edu, dict):
                        continue
                    institution = edu.get('institution', 'N/A')
                    doc.add_heading(institution, level=2)
                    
                    degree = edu.get('degree', '')
                    year = edu.get('graduation_year', '')
                    degree_text = degree
                    if year:
                        degree_text += f" ({year})"
                    doc.add_paragraph(degree_text)
            
            # Skills
            if resume.skills:
                doc.add_heading('Skills', level=1)
                doc.add_paragraph(", ".join(resume.skills))
            
            # Ensure directory exists
            os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
            
            doc.save(output_path)
            logger.info(f"Successfully exported DOCX to: {output_path}")
            return output_path
        except Exception as e:
            logger.error(f"Failed to export DOCX: {str(e)}", exc_info=True)
            raise

    @staticmethod
    def to_pdf(resume: ResumeData, output_path: str) -> str:
        """Export resume to PDF format."""
        try:
            if not resume.name:
                raise ValueError("Resume name is required")
            
            pdf = FPDF()
            pdf.add_page()
            font = _register_pdf_font(pdf)
            unicode_font = font == "DejaVu"
            bold = "B" if font == "Helvetica" or os.path.exists(_DEJAVU_BOLD) else ""
            italic = "I" if font == "Helvetica" or os.path.exists(_DEJAVU_OBLIQUE) else ""

            def _line(style: str, size: int, text: object, height: int = 8, align: str = "L") -> None:
                rendered = pdf_text(text, unicode_font).strip()
                if not rendered:
                    return
                pdf.set_x(pdf.l_margin)
                pdf.set_font(font, style, size)
                pdf.multi_cell(0, height, rendered, align=align)
            
            # Name
            _line(bold, 16, resume.name, 10, align="C")
            
            # Contact information
            contact_parts = [resume.email]
            if resume.phone:
                contact_parts.append(resume.phone)
            if resume.linkedin:
                contact_parts.append(resume.linkedin)
            
            if contact_parts:
                _line(italic, 10, " | ".join(contact_parts), 8, align="C")
            
            # Summary
            if resume.summary:
                _line(bold, 12, "Summary", 10)
                _line("", 10, resume.summary, 5)
            
            # Experience
            if resume.experience:
                _line(bold, 12, "Experience", 10)
                for exp in resume.experience:
                    if not isinstance(exp, dict):
                        continue
                    role = exp.get('role', 'N/A')
                    company = exp.get('company', 'N/A')
                    _line(bold, 10, f"{role} at {company}", 7)
                    
                    if exp.get('duration'):
                        _line(italic, 9, exp['duration'], 7)
                    
                    if exp.get('description'):
                        _line("", 10, exp['description'], 5)
            
            # Education
            if resume.education:
                _line(bold, 12, "Education", 10)
                for edu in resume.education:
                    if not isinstance(edu, dict):
                        continue
                    institution = edu.get('institution', 'N/A')
                    degree = edu.get('degree', '')
                    year = edu.get('graduation_year', '')
                    
                    _line(bold, 10, institution, 7)
                    degree_text = degree
                    if year:
                        degree_text += f" ({year})"
                    _line("", 10, degree_text, 7)
            
            # Skills
            if resume.skills:
                _line(bold, 12, "Skills", 10)
                skill_values = [str(skill) for skill in resume.skills if skill]
                _line("", 10, ", ".join(skill_values), 5)
            
            # Ensure directory exists
            os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
            
            pdf.output(output_path)
            logger.info(f"Successfully exported PDF to: {output_path}")
            return output_path
        except Exception as e:
            logger.error(f"Failed to export PDF: {str(e)}", exc_info=True)
            raise

    @staticmethod
    def to_latex(resume: ResumeData) -> str:
        """Export resume to LaTeX format."""
        try:
            if not resume.name:
                raise ValueError("Resume name is required")
            
            name_parts = resume.name.split()
            first_name = escape_latex(name_parts[0]) if name_parts else ""
            last_name = escape_latex(" ".join(name_parts[1:])) if len(name_parts) > 1 else ""
            
            latex = r"""
\documentclass[11pt,a4paper,sans]{moderncv}
\moderncvstyle{casual}
\moderncvcolor{blue}
\name{""" + first_name + "}{" + last_name + r"""}
\address{""" + escape_latex(resume.email) + r"""}
"""
            
            if resume.phone:
                latex += "\\phone[mobile]{" + escape_latex(resume.phone) + "}\n"
            
            if resume.linkedin:
                latex += "\\social[linkedin]{" + escape_latex(resume.linkedin) + "}\n"
            
            latex += r"""
\begin{document}
\makecvtitle

"""
            
            # Summary
            if resume.summary:
                latex += r"\section{Summary}" + "\n"
                latex += escape_latex(resume.summary) + "\n\n"
            
            # Experience
            if resume.experience:
                latex += r"\section{Experience}" + "\n"
                for exp in resume.experience:
                    if not isinstance(exp, dict):
                        continue
                    duration = escape_latex(exp.get('duration', ''))
                    role = escape_latex(exp.get('role', ''))
                    company = escape_latex(exp.get('company', ''))
                    description = escape_latex(exp.get('description', ''))
                    latex += f"\\cventry{{{duration}}}{{{role}}}{{{company}}}{{}}{{}}{{{description}}}\n"
                latex += "\n"
            
            # Education
            if resume.education:
                latex += r"\section{Education}" + "\n"
                for edu in resume.education:
                    if not isinstance(edu, dict):
                        continue
                    year = escape_latex(str(edu.get('graduation_year', '')))
                    degree = escape_latex(edu.get('degree', ''))
                    institution = escape_latex(edu.get('institution', ''))
                    latex += f"\\cventry{{{year}}}{{{degree}}}{{{institution}}}{{}}{{}}{{}}\n"
                latex += "\n"
            
            # Skills
            if resume.skills:
                latex += r"\section{Skills}" + "\n"
                skills_text = ", ".join(escape_latex(skill) for skill in resume.skills if skill)
                latex += f"\\cvitem{{Core Skills}}{{{skills_text}}}\n\n"
            
            latex += r"""
\end{document}
"""
            
            logger.info("Successfully generated LaTeX export")
            return latex
        except Exception as e:
            logger.error(f"Failed to export LaTeX: {str(e)}", exc_info=True)
            raise

